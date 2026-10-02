import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_document_from_md(md_file_path, output_docx_path):
    with open(md_file_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = docx.Document()

    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(30, 41, 59) # Slate 800

    in_table = False
    table_rows = []
    in_code_block = False
    code_lines = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        
        # Clean rows
        parsed_rows = []
        for r in table_rows:
            cells = [c.strip() for c in r.strip('|').split('|')]
            if any('---' in c for c in cells):
                continue
            parsed_rows.append(cells)

        if not parsed_rows:
            in_table = False
            table_rows = []
            return

        num_cols = max(len(r) for r in parsed_rows)
        tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        for r_idx, row in enumerate(parsed_rows):
            for c_idx in range(num_cols):
                val = row[c_idx] if c_idx < len(row) else ""
                cell = tbl.cell(r_idx, c_idx)
                cell.text = val
                set_cell_margins(cell, 120, 120, 160, 160)
                
                # Format cell text
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                
                if r_idx == 0:
                    set_cell_background(cell, "1E293B") # Dark slate
                    for run in p.runs:
                        run.font.bold = True
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                    set_cell_background(cell, bg_color)
                    for run in p.runs:
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = RGBColor(51, 65, 85)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        in_table = False
        table_rows = []

    def flush_code():
        nonlocal in_code_block, code_lines
        if not code_lines:
            in_code_block = False
            return
        code_text = "".join(code_lines)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.left_indent = Inches(0.2)
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(30, 41, 59)
        in_code_block = False
        code_lines = []

    for line in lines:
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith('```'):
            if in_code_block:
                flush_code()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Handle tables
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_rows.append(stripped)
            continue
        elif in_table:
            flush_table()

        # Handle headings
        if stripped.startswith('# '):
            h = doc.add_heading(level=0)
            run = h.add_run(stripped[2:])
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            continue
        elif stripped.startswith('## '):
            h = doc.add_heading(level=1)
            run = h.add_run(stripped[3:])
            run.font.name = 'Calibri'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138) # Corporate Navy
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            continue
        elif stripped.startswith('### '):
            h = doc.add_heading(level=2)
            run = h.add_run(stripped[4:])
            run.font.name = 'Calibri'
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(2)
            continue

        # Horizontal rule
        if stripped == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run('_______________________________________________________________________________')
            run.font.color.rgb = RGBColor(203, 213, 225)
            run.font.size = Pt(8)
            continue

        # Empty lines
        if not stripped:
            continue

        # Bullet list items
        if stripped.startswith('* ') or stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            clean_text = stripped[2:]
            # format bold
            parts = re.split(r'(\*\*.*?\*\*)', clean_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                else:
                    p.add_run(part)
            continue

        # Numbered list items
        num_match = re.match(r'^(\d+\.)\s+(.*)', stripped)
        if num_match:
            num_prefix = num_match.group(1)
            clean_text = num_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run_num = p.add_run(num_prefix + " ")
            run_num.font.bold = True
            parts = re.split(r'(\*\*.*?\*\*)', clean_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                else:
                    p.add_run(part)
            continue

        # Regular paragraphs
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        parts = re.split(r'(\*\*.*?\*\*)', stripped)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.font.bold = True
            else:
                p.add_run(part)

    if in_table:
        flush_table()
    if in_code_block:
        flush_code()

    doc.save(output_docx_path)
    print(f"Successfully generated: {output_docx_path}")

if __name__ == '__main__':
    md_path = r'c:\Users\shara\OneDrive\Desktop\Hitesh Projects\Logistics B2B and 2PL\DOCUMENTATION.md'
    docx_path = r'c:\Users\shara\OneDrive\Desktop\Hitesh Projects\Logistics B2B and 2PL\DOCUMENTATION.docx'
    create_document_from_md(md_path, docx_path)

    desktop_docx = r'c:\Users\shara\OneDrive\Desktop\DOCUMENTATION.docx'
    try:
        import shutil
        shutil.copy2(docx_path, desktop_docx)
        print(f"Successfully synced to desktop: {desktop_docx}")
    except Exception as e:
        print(f"Desktop copy notice: {e}")
